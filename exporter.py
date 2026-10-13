"""Harajatlarni Excel (.xlsx) va CSV formatiga eksport qilish."""
import csv
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from categories import CAT_BY_KEY

HEADERS = ["№", "Sana", "Kategoriya", "Summa (so'm)", "Izoh", "Kim kiritgan"]


def _cat_name(key: str) -> str:
    return CAT_BY_KEY.get(key, ("", "Boshqa"))[1]


def build_csv(rows: list) -> bytes:
    """CSV (UTF-8 BOM, ';' ajratgich) - Windows'dagi Excel to'g'ri ochishi uchun."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(HEADERS)
    for i, r in enumerate(rows, 1):
        writer.writerow([
            i,
            date.fromisoformat(r["spent_on"]).strftime("%d.%m.%Y"),
            _cat_name(r["category"]),
            r["amount"],
            r["note"],
            r["user_name"],
        ])
    return buf.getvalue().encode("utf-8-sig")


def _style_header(ws, columns: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for col in range(1, columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _set_widths(ws, widths: list) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(rows: list) -> bytes:
    """Excel fayl: 'Harajatlar', 'Kategoriyalar' va 'Oylar' varaqlari."""
    wb = Workbook()

    # --- 1-varaq: barcha harajatlar
    ws = wb.active
    ws.title = "Harajatlar"
    ws.append(HEADERS)
    _style_header(ws, len(HEADERS))
    total = 0
    for i, r in enumerate(rows, 1):
        ws.append([
            i,
            date.fromisoformat(r["spent_on"]),
            _cat_name(r["category"]),
            int(r["amount"]),
            r["note"],
            r["user_name"],
        ])
        total += int(r["amount"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[1].number_format = "DD.MM.YYYY"
        row[3].number_format = "#,##0"
    ws.append(["", "", "JAMI:", total, "", ""])
    last = ws.max_row
    ws.cell(row=last, column=3).font = Font(bold=True)
    ws.cell(row=last, column=4).font = Font(bold=True)
    ws.cell(row=last, column=4).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{max(last - 1, 1)}"
    _set_widths(ws, [6, 13, 38, 16, 36, 20])

    # --- 2-varaq: kategoriyalar bo'yicha
    by_cat = defaultdict(lambda: [0, 0])
    for r in rows:
        by_cat[r["category"]][0] += 1
        by_cat[r["category"]][1] += int(r["amount"])
    ws2 = wb.create_sheet("Kategoriyalar")
    ws2.append(["Kategoriya", "Yozuvlar soni", "Summa (so'm)", "Ulushi"])
    _style_header(ws2, 4)
    for key, (cnt, amount) in sorted(by_cat.items(), key=lambda kv: -kv[1][1]):
        ws2.append([_cat_name(key), cnt, amount, (amount / total) if total else 0])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        row[2].number_format = "#,##0"
        row[3].number_format = "0.0%"
    _set_widths(ws2, [38, 15, 16, 10])

    # --- 3-varaq: oylar bo'yicha
    by_month = defaultdict(lambda: [0, 0])
    for r in rows:
        month = r["spent_on"][:7]
        by_month[month][0] += 1
        by_month[month][1] += int(r["amount"])
    ws3 = wb.create_sheet("Oylar")
    ws3.append(["Oy", "Yozuvlar soni", "Summa (so'm)"])
    _style_header(ws3, 3)
    for month, (cnt, amount) in sorted(by_month.items()):
        ws3.append([month, cnt, amount])
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        row[2].number_format = "#,##0"
    _set_widths(ws3, [12, 15, 16])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
